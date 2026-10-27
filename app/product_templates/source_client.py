"""Trusted bounded CSV/JSON/XML/ZIP/XLSX readers, usable by app and Airflow DAGs."""
import csv
import io
import json
import os
from pathlib import PurePosixPath
import urllib.request
import zipfile

MAX_BYTES=8_000_000

def download(source_id):
    import re
    if not re.fullmatch(r'[a-z][a-z0-9_-]{0,39}',source_id):raise ValueError('Invalid source ID')
    with urllib.request.urlopen(os.environ['SOURCE_GATEWAY_URL']+'/sources/'+source_id,timeout=30) as response:
        data=response.read(MAX_BYTES+1)
    if len(data)>MAX_BYTES:raise ValueError('Source too large')
    return data


def parse(data,format,*,encoding='utf-8-sig'):
    if len(data)>MAX_BYTES:raise ValueError('Source too large')
    if format=='csv':return list(csv.DictReader(io.StringIO(data.decode(encoding))))
    if format=='json':return json.loads(data)
    if format=='xml':
        from lxml import etree
        if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():raise ValueError('DTD/entities forbidden')
        return etree.fromstring(data,parser=etree.XMLParser(resolve_entities=False,no_network=True,load_dtd=False,huge_tree=False))
    if format in {'zip','xlsx'}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos=archive.infolist()
            if len(infos)>200 or sum(i.file_size for i in infos)>32_000_000:raise ValueError('Expanded source too large')
            for item in infos:
                path=PurePosixPath(item.filename)
                if path.is_absolute() or '..' in path.parts or '\\' in item.filename or item.flag_bits&1:raise ValueError('Unsafe ZIP path/encryption')
            if format=='zip':return {i.filename:archive.read(i) for i in infos if not i.is_dir()}
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        try:
            result={};count=0
            for sheet in workbook:
                rows=[]
                for row in sheet.iter_rows(values_only=True):
                    count+=len(row)
                    if count>200000:raise ValueError('Workbook too large')
                    rows.append(list(row))
                result[sheet.title]=rows
            return result
        finally:workbook.close()
    raise ValueError('Unsupported format')
